#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent.parent

REQUIRED_PDFS = [
    "chiffrement-reseau_2025-2026_etudiants_tp2.pdf",
    "cryptography_base_course_student.pdf",
    "cryptography_communication_course_student.pdf",
]

REQUIRED_REPORTS = [
    "compte_rendu_tp2_chiffrement_reseau.md",
    "compte_rendu_tp2_chiffrement_reseau.docx",
]


def count_pdf_pages(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def count_docx_paragraphs(path: Path) -> int:
    return len(Document(path).paragraphs)


def main() -> int:
    print("Validation de l'environnement du depot")
    print(f"Racine: {ROOT}")
    print("")

    missing = []
    for name in REQUIRED_PDFS + REQUIRED_REPORTS:
        path = ROOT / name
        if not path.exists():
            missing.append(name)

    if missing:
        print("Fichiers manquants:")
        for name in missing:
            print(f"- {name}")
        return 1

    print("Fichiers obligatoires presents : OK")
    print("")

    pdf_summaries = []
    for name in REQUIRED_PDFS:
        path = ROOT / name
        pdf_summaries.append((name, count_pdf_pages(path), path.stat().st_size))

    print("Verification PDF")
    for name, pages, size in pdf_summaries:
        print(f"- {name}: {pages} pages, {size} octets")

    md_path = ROOT / "compte_rendu_tp2_chiffrement_reseau.md"
    md_text = md_path.read_text(encoding="utf-8")
    md_lines = len(md_text.splitlines())
    md_has_tp = "FTPS" in md_text and "HTTPS" in md_text

    docx_path = ROOT / "compte_rendu_tp2_chiffrement_reseau.docx"
    doc = Document(docx_path)
    docx_paragraphs = count_docx_paragraphs(docx_path)
    first_paragraph = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), "")

    print("")
    print("Verification du rendu")
    print(f"- Markdown: {md_lines} lignes")
    print(f"- Sections FTPS/HTTPS detectees: {'oui' if md_has_tp else 'non'}")
    print(f"- DOCX: {docx_paragraphs} paragraphes")
    print(f"- Premier paragraphe DOCX: {first_paragraph}")

    if not md_has_tp:
        print("")
        print("Le rendu ne semble pas contenir les sections attendues sur FTPS et HTTPS.")
        return 1

    print("")
    print("Demonstration reussie : l'environnement Python est operationnel")
    print("et les livrables principaux du depot sont lisibles et valides.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
