from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


INPUT_PATH = Path("/workspace/compte_rendu_tp2_chiffrement_reseau_propose.md")
OUTPUT_PATH = Path("/workspace/compte_rendu_tp2_chiffrement_reseau_propose.docx")


def flush_paragraph(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = " ".join(line.strip() for line in buffer).strip()
    if text:
        doc.add_paragraph(text)
    buffer.clear()


def main() -> None:
    lines = INPUT_PATH.read_text(encoding="utf-8").splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    paragraph_buffer: list[str] = []
    first_title = True

    for line in lines:
        stripped = line.strip()

        if not stripped:
            flush_paragraph(doc, paragraph_buffer)
            continue

        if stripped == "---":
            flush_paragraph(doc, paragraph_buffer)
            if not first_title:
                doc.add_page_break()
            continue

        if stripped.startswith("# "):
            flush_paragraph(doc, paragraph_buffer)
            if first_title:
                doc.add_heading(stripped[2:].strip(), level=0)
                first_title = False
            else:
                doc.add_heading(stripped[2:].strip(), level=1)
            continue

        if stripped.startswith("## "):
            flush_paragraph(doc, paragraph_buffer)
            doc.add_heading(stripped[3:].strip(), level=1)
            continue

        if stripped.startswith("### "):
            flush_paragraph(doc, paragraph_buffer)
            doc.add_heading(stripped[4:].strip(), level=2)
            continue

        if stripped.startswith("- "):
            flush_paragraph(doc, paragraph_buffer)
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue

        if stripped.startswith("**") and stripped.endswith("**"):
            flush_paragraph(doc, paragraph_buffer)
            p = doc.add_paragraph()
            run = p.add_run(stripped.strip("*"))
            run.bold = True
            continue

        if "**" in stripped and stripped.endswith("  "):
            flush_paragraph(doc, paragraph_buffer)
            p = doc.add_paragraph()
            segments = stripped.split("**")
            for index, segment in enumerate(segments):
                if not segment:
                    continue
                run = p.add_run(segment.strip())
                if index % 2 == 1:
                    run.bold = True
                if index < len(segments) - 1:
                    run.add_break(WD_BREAK.LINE)
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph(doc, paragraph_buffer)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
